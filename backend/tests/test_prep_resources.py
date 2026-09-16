# 工单编号：人工智能NLP-Agent数字人项目-教育智能体-智能备课任务(17)
"""课程资源检索测试：上传 docx → 解析 → 混合检索 → 引用标注。"""
import io
from pathlib import Path

import pytest
from docx import Document
from fastapi import UploadFile

from app.db import Base
from app.models.prep import Course, CourseFile
from app.services.prep_resources import add_resource, list_resources, search_resources


@pytest.fixture
def db(tmp_path):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    engine = create_engine(f"sqlite:///{tmp_path / 'res.db'}")
    Base.metadata.create_all(engine)
    Session = sessionmaker(bind=engine)
    session = Session()
    session.add(Course(name="人工智能导论", subject="人工智能", owner_id=1))
    session.commit()
    yield session, tmp_path
    Base.metadata.drop_all(engine)


@pytest.fixture
def docx_path(tmp_path):
    """含中文段落与表格的课程资源样例。"""
    doc = Document()
    doc.add_paragraph("梯度下降是机器学习中最基础的优化算法，通过沿负梯度方向迭代更新参数。")
    doc.add_paragraph("线性回归通过拟合直线描述特征与目标值的关系，常用于房价预测。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "算法"; table.cell(0, 1).text = "场景"
    table.cell(1, 0).text = "线性回归"; table.cell(1, 1).text = "房价预测"
    path = tmp_path / "讲义.docx"
    doc.save(path)
    return path


class FakeEmbedder:
    """3 维固定向量：向量半程同分，检索结果由 BM25 半程主导，便于断言。"""
    dim = 3

    def embed_texts(self, texts):
        return [[1.0, 0.0, 0.0] for _ in texts]

    def embed_query(self, text):
        return [1.0, 0.0, 0.0]


def _add(db_pair, docx_path, filename="讲义.docx"):
    session, tmp_path = db_pair
    with open(docx_path, "rb") as f:
        file = UploadFile(filename=filename, file=io.BytesIO(f.read()))
        return add_resource(course_id=1, file=file, upload_dir=tmp_path / "uploads", db=session)


def test_add_resource_registers_course_file(db, docx_path):
    session, _ = db
    record = _add(db, docx_path)
    assert record.id > 0
    assert session.query(CourseFile).filter(CourseFile.course_id == 1).count() == 1
    listed = list_resources(1, session)
    assert listed[0].filename == "讲义.docx"


def test_search_resources_hits_and_ref_format(db, docx_path):
    session, tmp_path = db
    _add(db, docx_path)
    hits = search_resources(1, "梯度下降的优化原理是什么", tmp_path / "uploads", session,
                            top_k=3, embedder=FakeEmbedder())
    assert hits, "BM25 应命中语料"
    first = hits[0]
    assert first["chunk"].source == "讲义.docx"
    assert "梯度下降" in first["chunk"].text
    assert first["ref"].startswith("[1]")
    assert "来源：讲义.docx" in first["ref"]           # 引用标注格式
    assert first["score"] > 0


def test_search_resources_empty_course_returns_empty(db):
    session, tmp_path = db
    assert search_resources(1, "梯度下降", tmp_path / "uploads", session, embedder=FakeEmbedder()) == []
