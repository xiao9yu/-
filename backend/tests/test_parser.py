# 工单编号：人工智能NLP-Agent数字人项目-教育智能体-公共底座(16-20)
import io
from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook
from PIL import Image, ImageDraw
from pptx import Presentation
from pymupdf import open as open_pdf

from app.services.parser import parse_document
from app.services.parser.chunk import split_text


@pytest.fixture
def docx_with_table_image(tmp_path) -> Path:
    """生成包含段落、表格、图片的 docx 样例。"""
    doc = Document()
    doc.add_paragraph("机器学习是人工智能的一个分支。")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "模型"; table.cell(0, 1).text = "准确率"
    table.cell(1, 0).text = "决策树"; table.cell(1, 1).text = "0.85"
    # 生成一张图片插入
    img = Image.new("RGB", (80, 40), "white")
    ImageDraw.Draw(img).text((5, 10), "ML")
    buf = io.BytesIO(); img.save(buf, format="PNG"); buf.seek(0)
    doc.add_picture(buf)
    path = tmp_path / "讲义.docx"
    doc.save(path)
    return path


@pytest.fixture
def pptx_sample(tmp_path) -> Path:
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = "人工智能导论"
    path = tmp_path / "课件.pptx"
    prs.save(path)
    return path


@pytest.fixture
def xlsx_sample(tmp_path) -> Path:
    wb = Workbook()
    ws = wb.active
    ws.append(["知识点", "掌握人数"])
    ws.append(["梯度下降", 30])
    path = tmp_path / "成绩.xlsx"
    wb.save(path)
    return path


@pytest.fixture
def pdf_sample(tmp_path) -> Path:
    path = tmp_path / "教材.pdf"
    doc = open_pdf()
    page = doc.new_page()
    page.insert_text((72, 72), "神经网络由多层神经元组成。", fontname="china-s")
    doc.save(path); doc.close()
    return path


def test_split_text_with_overlap():
    parts = split_text("A" * 30, chunk_size=10, overlap=2)
    assert len(parts) == 4  # 10+10+10+6（最后一片不足一个整块）
    assert parts[0] == "A" * 10


def test_parse_docx_kinds(docx_with_table_image):
    chunks = parse_document(docx_with_table_image)
    kinds = {c.kind for c in chunks}
    assert "text" in kinds and "table" in kinds and "image" in kinds
    table_chunk = next(c for c in chunks if c.kind == "table")
    assert "模型" in table_chunk.text and "决策树" in table_chunk.text


def test_parse_pptx_text(pptx_sample):
    chunks = parse_document(pptx_sample)
    assert any("人工智能导论" in c.text for c in chunks)


def test_parse_xlsx_table(xlsx_sample):
    chunks = parse_document(xlsx_sample)
    assert any(c.kind == "table" and "梯度下降" in c.text for c in chunks)


def test_parse_pdf_text(pdf_sample):
    chunks = parse_document(pdf_sample)
    assert any("神经网络" in c.text for c in chunks)
    assert all(c.page == 1 for c in chunks)


def test_parse_image_ocr(monkeypatch, tmp_path):
    from app.services.parser import image_parser

    img = Image.new("RGB", (80, 40), "white")
    path = tmp_path / "photo.png"
    img.save(path)
    monkeypatch.setattr(image_parser, "ocr_image", lambda p: "识别出的板书文字")
    chunks = parse_document(path)
    assert len(chunks) == 1
    assert chunks[0].kind == "image"
    assert chunks[0].text == "识别出的板书文字"
    assert chunks[0].meta["image_path"] == str(path)


def test_parse_unsupported_returns_empty(tmp_path):
    path = tmp_path / "x.xyz"
    path.write_text("abc")
    assert parse_document(path) == []
