# 工单编号：人工智能NLP-Agent数字人项目-教育智能体-智能备课任务(17)
"""备课 API 测试：课程 CRUD/资源/生成/版本/导出 + 权限拦截（真实临时库 + TestClient）。"""
import io

import pytest
from docx import Document
from fastapi.testclient import TestClient

from app.db import Base, get_db
from app.main import app


@pytest.fixture
def db(tmp_path):
    from sqlalchemy import create_engine
    from sqlalchemy.orm import sessionmaker
    engine = create_engine(f"sqlite:///{tmp_path / 'api.db'}", connect_args={"check_same_thread": False})
    TestingSession = sessionmaker(bind=engine, autoflush=False, autocommit=False)
    Base.metadata.create_all(engine)

    def override_get_db():
        session = TestingSession()
        try:
            yield session
        finally:
            session.close()

    app.dependency_overrides[get_db] = override_get_db
    yield TestingSession()
    app.dependency_overrides.clear()
    Base.metadata.drop_all(engine)


@pytest.fixture
def client(db):
    with TestClient(app) as c:
        yield c


def _register_login(client, username, role):
    client.post("/api/auth/register", json={"username": username, "password": "pass123456",
                                            "role": role, "real_name": "测试"})
    resp = client.post("/api/auth/login", json={"username": username, "password": "pass123456"})
    return {"Authorization": f"Bearer {resp.json()['access_token']}"}


def _create_course(client, headers, name="人工智能导论"):
    resp = client.post("/api/prep/courses", json={"name": name, "subject": "人工智能",
                                                  "description": "测试课程"}, headers=headers)
    assert resp.status_code == 200
    return resp.json()


def test_student_cannot_create_course(client):
    headers = _register_login(client, "stu_prep", "student")
    resp = client.post("/api/prep/courses", json={"name": "x", "subject": "y"}, headers=headers)
    assert resp.status_code == 403


def test_teacher_course_crud_and_collaborator(client):
    headers = _register_login(client, "t1", "teacher")
    course = _create_course(client, headers)
    listed = client.get("/api/prep/courses", headers=headers).json()
    assert any(c["id"] == course["id"] for c in listed)
    # 非成员不可见
    other = _register_login(client, "t2", "teacher")
    assert client.get(f"/api/prep/courses/{course['id']}", headers=other).status_code == 403
    # 添加协作者后可见
    resp = client.post("/api/auth/login", json={"username": "t2", "password": "pass123456"})
    uid2 = resp.json()["user"]["id"]
    assert client.post(f"/api/prep/courses/{course['id']}/collaborators",
                       json={"user_id": uid2}, headers=headers).status_code == 200
    assert client.get(f"/api/prep/courses/{course['id']}", headers=other).status_code == 200


def test_generate_endpoint_mocks_gateway_and_uses_resources(client, db, tmp_path, monkeypatch):
    from app.config import settings
    from app.services import prep_generator
    from app.services import prep_resources
    from app.services.prep_resources import search_resources
    # 上传目录与检索目录统一指向测试临时目录（API 端点内用 settings.upload_dir）
    monkeypatch.setattr(settings, "upload_dir", str(tmp_path / "uploads"))
    headers = _register_login(client, "t3", "teacher")
    course = _create_course(client, headers)
    # 上传课程资源（真实 docx）
    doc = Document()
    doc.add_paragraph("梯度下降通过沿负梯度方向迭代更新参数逼近最优解。")
    doc.save(tmp_path / "讲义.docx")
    with open(tmp_path / "讲义.docx", "rb") as f:
        client.post(f"/api/prep/courses/{course['id']}/resources",
                    files={"file": ("讲义.docx", f, "application/octet-stream")}, headers=headers)
    # 检索资源
    hits = search_resources(course["id"], "梯度下降", settings.upload_dir, db,
                            embedder=_FakeEmb())
    assert hits and "讲义.docx" in hits[0]["ref"]
    # mock 生成网关
    fake = _FakeLLM({"标题": "教案", "教学目标": ["理解梯度下降"], "教学重点": ["迭代"],
                     "教学难点": ["学习率"], "教学过程": [{"环节": "导入", "内容": "回顾", "时长": "5分钟"}],
                     "作业": "习题", "板书设计": "公式"})
    monkeypatch.setattr(prep_generator, "get_gateway", lambda: fake)
    # mock 嵌入模型：API 路径内部检索默认用真实 bge-m3，本机无法访问 HuggingFace 会卡死；
    # 按"生成接口测试 mock LLM/检索"约定改用假嵌入（与直接检索断言用同一 _FakeEmb）
    monkeypatch.setattr(prep_resources, "get_embedder", lambda: _FakeEmb())
    resp = client.post(f"/api/prep/courses/{course['id']}/generate",
                       json={"type": "plan", "chapter": "第3章", "objectives": "理解梯度下降",
                             "hours": "2课时", "query": "梯度下降"}, headers=headers)
    assert resp.status_code == 200
    data = resp.json()
    assert data["content"]["标题"] == "教案"
    assert data["citations"], "query 非空时应检索资源并返回引用"
    assert "梯度下降" in fake.calls[0]["messages"][-1]["content"]  # prompt 含课程信息


def test_lesson_save_version_restore_and_citations(client):
    headers = _register_login(client, "t4", "teacher")
    course = _create_course(client, headers)
    content_v1 = {"标题": "第一版教案", "教学目标": ["目标A"]}
    resp = client.post(f"/api/prep/courses/{course['id']}/lessons",
                       json={"title": "教案1", "lesson_type": "plan", "content_json": content_v1},
                       headers=headers)
    assert resp.status_code == 200
    lesson = resp.json()
    assert lesson["version"] == 1
    # 保存新版本
    resp = client.put(f"/api/prep/lessons/{lesson['id']}",
                      json={"content_json": {"标题": "第二版教案", "教学目标": ["目标B"]}},
                      headers=headers)
    assert resp.status_code == 200
    assert resp.json()["version"] == 2
    versions = client.get(f"/api/prep/lessons/{lesson['id']}/versions", headers=headers).json()
    assert len(versions) == 2
    # 恢复 v1
    resp = client.post(f"/api/prep/lessons/{lesson['id']}/restore", json={"version": 1}, headers=headers)
    assert resp.status_code == 200
    assert resp.json()["content_json"]["标题"] == "第一版教案"
    assert resp.json()["version"] == 3
    # 保存引用
    resp = client.post(f"/api/prep/lessons/{lesson['id']}/citations",
                       json={"ref_no": 1, "source": "教材.pdf", "page": 3, "excerpt": "梯度下降"},
                       headers=headers)
    assert resp.status_code == 200


def test_media_registration(client):
    """多媒体登记：文件先走底座 files 表，再关联到教案（media_files）。"""
    headers = _register_login(client, "t6", "teacher")
    course = _create_course(client, headers)
    lesson = client.post(f"/api/prep/courses/{course['id']}/lessons",
                         json={"title": "教案", "lesson_type": "plan",
                               "content_json": {"标题": "教案"}},
                         headers=headers).json()
    # 上传图片（底座文件服务，端点为 /api/files/upload）
    import io
    png_bytes = b"\x89PNG\r\n\x1a\n" + b"0" * 32
    resp = client.post("/api/files/upload", files={"file": ("图片.png", io.BytesIO(png_bytes), "image/png")},
                       headers=headers)
    assert resp.status_code == 200
    file_id = resp.json()["id"]
    resp = client.post(f"/api/prep/lessons/{lesson['id']}/media",
                       json={"file_id": file_id}, headers=headers)
    assert resp.status_code == 200
    assert resp.json()["ok"] is True


def test_export_endpoint_docx(client, tmp_path):
    headers = _register_login(client, "t5", "teacher")
    course = _create_course(client, headers)
    lesson = client.post(f"/api/prep/courses/{course['id']}/lessons",
                         json={"title": "教案", "lesson_type": "plan",
                               "content_json": {"标题": "梯度下降教案", "教学目标": ["理解梯度下降"],
                                                "教学重点": ["迭代"], "教学难点": ["学习率"],
                                                "教学过程": [{"环节": "导入", "内容": "回顾", "时长": "5分钟"}],
                                                "作业": "习题", "板书设计": "公式"}},
                         headers=headers).json()
    resp = client.get(f"/api/prep/lessons/{lesson['id']}/export?format=docx", headers=headers)
    assert resp.status_code == 200
    assert resp.headers["content-type"].startswith("application/octet-stream") or "docx" in resp.headers["content-disposition"]
    from docx import Document as Doc
    out = tmp_path / "export.docx"
    out.write_bytes(resp.content)
    text = "\n".join(p.text for p in Doc(str(out)).paragraphs)
    assert "梯度下降教案" in text


class _FakeEmb:
    dim = 3

    def embed_texts(self, texts):
        return [[1.0, 0.0, 0.0] for _ in texts]

    def embed_query(self, text):
        return [1.0, 0.0, 0.0]


class _FakeLLM:
    def __init__(self, result):
        self.result = result
        self.calls = []

    def chat_json(self, messages, temperature=0.3):
        self.calls.append({"messages": messages, "temperature": temperature})
        return self.result
