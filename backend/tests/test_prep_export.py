# 工单编号：人工智能NLP-Agent数字人项目-教育智能体-智能备课任务(17)
"""导出服务测试：教案→docx、课件→pptx、习题→pdf，生成后读回断言内容。"""
from pathlib import Path

import pytest

from app.services.prep_export import (
    export_case_docx,
    export_courseware_pptx,
    export_exercises_pdf,
    export_lesson_docx,
)

LESSON_PLAN = {
    "标题": "梯度下降教案",
    "教学目标": ["理解梯度下降原理"],
    "教学重点": ["迭代公式"],
    "教学难点": ["学习率选择"],
    "教学过程": [{"环节": "导入", "内容": "回顾线性回归", "时长": "5分钟"}],
    "作业": "完成课后习题",
    "板书设计": "迭代公式示意图",
}

COURSEWARE = {
    "标题": "机器学习基础课件",
    "幻灯片": [{"标题": "梯度下降", "要点": ["沿负梯度方向迭代", "学习率控制步长"]}],
}

EXERCISES = [
    {"题干": "梯度下降中控制步长的参数是", "选项": ["A.学习率", "B.批量大小"],
     "答案": "A", "解析": "学习率控制步长", "知识点": "梯度下降", "难度": "易"},
]

CASE = {
    "标题": "垃圾邮件分类教学案例",
    "案例背景": "某电商平台每日收到大量垃圾邮件。",
    "案例描述": "使用朴素贝叶斯对邮件进行二分类。",
    "问题": "如何设计特征并评估模型效果？",
    "案例分析": "提取词频特征，交叉验证评估。",
    "结论": "垃圾邮件召回率达 95%。",
}


def test_export_lesson_docx(tmp_path):
    out = export_lesson_docx(LESSON_PLAN, tmp_path / "教案.docx")
    assert out.exists()
    from docx import Document
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "梯度下降教案" in text and "理解梯度下降原理" in text


def test_export_courseware_pptx(tmp_path):
    out = export_courseware_pptx(COURSEWARE, tmp_path / "课件.pptx")
    assert out.exists()
    from pptx import Presentation
    prs = Presentation(str(out))
    assert len(prs.slides) == 1
    assert prs.slides[0].shapes.title.text == "梯度下降"


def test_export_exercises_pdf(tmp_path):
    out = export_exercises_pdf(EXERCISES, "诊断试题", tmp_path / "习题.pdf")
    assert out.exists()
    import pymupdf
    doc = pymupdf.open(str(out))
    text = doc[0].get_text()
    doc.close()
    assert "梯度下降" in text          # PDF 中文正常渲染（字体注册生效）
    assert "诊断试题" in text


def test_export_case_docx(tmp_path):
    out = export_case_docx(CASE, tmp_path / "案例.docx")
    assert out.exists()
    from docx import Document
    doc = Document(str(out))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "垃圾邮件分类教学案例" in text
    assert "朴素贝叶斯" in text and "95%" in text
