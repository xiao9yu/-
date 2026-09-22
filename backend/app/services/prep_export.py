# 工单编号：人工智能NLP-Agent数字人项目-教育智能体-智能备课任务(17)
"""导出服务：教案→Word（python-docx）、课件大纲→PPT（python-pptx）、习题/月考题→PDF（reportlab）。
内容以生成时的结构化数据为准（V1 范围声明见计划 Global Constraints）。
"""
from pathlib import Path

from docx import Document
from docx.shared import Pt
from pptx import Presentation
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.cidfonts import UnicodeCIDFont
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.pdfgen import canvas

_CN_FONT_NAME: str | None = None


def _register_cn_font() -> str:
    """注册中文字体：优先 Windows 黑体（simhei.ttf），找不到回退 reportlab 内置 CID 字体。
    返回注册名。"""
    global _CN_FONT_NAME
    if _CN_FONT_NAME:
        return _CN_FONT_NAME
    for path in (r"C:\Windows\Fonts\simhei.ttf", r"C:\Windows\Fonts\simfang.ttf"):
        try:
            if Path(path).exists():
                pdfmetrics.registerFont(TTFont("CnFont", path))
                _CN_FONT_NAME = "CnFont"
                return "CnFont"
        except Exception:
            continue
    pdfmetrics.registerFont(UnicodeCIDFont("STSong-Light"))
    _CN_FONT_NAME = "STSong-Light"
    return "STSong-Light"


def export_lesson_docx(content: dict, out_path: Path) -> Path:
    """教案 JSON → Word 文档。"""
    doc = Document()
    doc.add_heading(content.get("标题", "教案"), level=1)
    doc.add_heading("教学目标", level=2)
    for item in content.get("教学目标", []):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("教学重点", level=2)
    for item in content.get("教学重点", []):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("教学难点", level=2)
    for item in content.get("教学难点", []):
        doc.add_paragraph(item, style="List Bullet")
    doc.add_heading("教学过程", level=2)
    for step in content.get("教学过程", []):
        doc.add_paragraph(f"{step.get('环节', '')}（{step.get('时长', '')}）：{step.get('内容', '')}")
    doc.add_heading("作业", level=2)
    doc.add_paragraph(content.get("作业", ""))
    doc.add_heading("板书设计", level=2)
    doc.add_paragraph(content.get("板书设计", ""))
    for p in doc.paragraphs:
        p.style.font.size = Pt(11)
    doc.save(str(out_path))
    return out_path


def export_courseware_pptx(content: dict, out_path: Path) -> Path:
    """课件大纲 JSON → PPT 文档（标题+要点列表；每页讲稿写入备注栏供教师放映讲解）。"""
    prs = Presentation()
    for slide_data in content.get("幻灯片", []):
        slide = prs.slides.add_slide(prs.slide_layouts[1])  # 标题+内容版式
        slide.shapes.title.text = slide_data.get("标题", "")
        body = slide.placeholders[1].text_frame
        body.text = "\n".join(slide_data.get("要点", []))
        script = (slide_data.get("讲稿") or "").strip()
        if script:
            # python-pptx 无显式建备注 API：访问 notes_slide 时惰性创建 notesSlide 部件
            slide.notes_slide.notes_text_frame.text = script
    prs.save(str(out_path))
    return out_path


def export_exercises_pdf(exercises: list[dict], title: str, out_path: Path) -> Path:
    """习题列表 → PDF（选择题：题干/选项/答案/解析，中文渲染）。"""
    font = _register_cn_font()
    pdf = canvas.Canvas(str(out_path), pagesize=A4)
    pdf.setTitle(title)
    y = A4[1] - 2 * cm
    pdf.setFont(font, 18)
    pdf.drawCentredString(A4[0] / 2, y, title)
    y -= 1.2 * cm
    pdf.setFont(font, 11)
    for i, ex in enumerate(exercises, start=1):
        lines = [f"{i}. {ex.get('题干', '')}"]
        lines.extend(ex.get("选项", []))
        lines.append(f"答案：{ex.get('答案', '')}　解析：{ex.get('解析', '')}")
        lines.append(f"知识点：{ex.get('知识点', '')}　难度：{ex.get('难度', '')}")
        for line in lines:
            if y < 2 * cm:          # 翻页
                pdf.showPage()
                pdf.setFont(font, 11)
                y = A4[1] - 2 * cm
            pdf.drawString(2 * cm, y, line)
            y -= 0.7 * cm
        y -= 0.5 * cm
    pdf.save()
    return out_path


def export_case_docx(content: dict, out_path: Path) -> Path:
    """教学案例 JSON → Word 文档。"""
    doc = Document()
    doc.add_heading(content.get("标题", "教学案例"), level=1)
    for heading, key in [("案例背景", "案例背景"), ("案例描述", "案例描述"), ("问题", "问题"),
                         ("案例分析", "案例分析"), ("结论", "结论")]:
        doc.add_heading(heading, level=2)
        doc.add_paragraph(content.get(key, ""))
    for p in doc.paragraphs:
        p.style.font.size = Pt(11)
    doc.save(str(out_path))
    return out_path
