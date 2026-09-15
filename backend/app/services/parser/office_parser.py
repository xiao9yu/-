# 工单编号：人工智能NLP-Agent数字人项目-教育智能体-公共底座(16-20)
"""Office 解析：docx（段落/表格/图片保序）、pptx（幻灯片文本）、xlsx（表格化）。"""
import uuid
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from openpyxl import load_workbook
from pptx import Presentation

from .chunk import Chunk, split_text
from .image_parser import ocr_image


def _flush(buf: list[str], source: str, chunks: list[Chunk]):
    text = "\n".join(buf).strip()
    if text:
        for part in split_text(text):
            chunks.append(Chunk(text=part, kind="text", source=source))
    buf.clear()


def _table_to_markdown(table) -> str:
    rows = []
    for row in table.rows:
        cells = [c.text.replace("\n", " ").strip() for c in row.cells]
        rows.append("| " + " | ".join(cells) + " |")
    if not rows:
        return ""
    ncols = rows[0].count("|") - 1
    header_sep = "| " + " | ".join(["---"] * ncols) + " |"
    return "\n".join([rows[0], header_sep, *rows[1:]])


def parse_docx(path: Path, extract_dir: Path | None = None) -> list[Chunk]:
    doc = Document(str(path))
    chunks: list[Chunk] = []
    buf: list[str] = []
    for child in doc.element.body.iterchildren():
        tag = child.tag
        if tag == qn("w:p"):
            para = [t.text for t in child.iter(qn("w:t")) if t.text]
            has_pic = len(child.findall(".//" + qn("pic:pic"))) > 0
            if has_pic:
                _flush(buf, path.name, chunks)
                # 提取内嵌图片
                img_rid = None
                for blip in child.iter(qn("a:blip")):
                    img_rid = blip.get(qn("r:embed"))
                    break
                if img_rid and img_rid in doc.part.related_parts:
                    blob = doc.part.related_parts[img_rid].blob
                    img_path = (extract_dir or Path("uploads/extracted")) / f"{uuid.uuid4().hex}.png"
                    img_path.parent.mkdir(parents=True, exist_ok=True)
                    img_path.write_bytes(blob)
                    chunks.append(Chunk(
                        text=ocr_image(img_path),
                        kind="image",
                        source=path.name,
                        meta={"image_path": str(img_path)},
                    ))
            else:
                buf.append("".join(para))
        elif tag == qn("w:tbl"):
            _flush(buf, path.name, chunks)
            table = None
            # 用 python-docx 的 Table 包装当前 w:tbl 元素
            from docx.table import Table
            table = Table(child, doc)
            md = _table_to_markdown(table)
            if md:
                chunks.append(Chunk(text=md, kind="table", source=path.name))
    _flush(buf, path.name, chunks)
    return chunks


def parse_pptx(path: Path) -> list[Chunk]:
    prs = Presentation(str(path))
    chunks: list[Chunk] = []
    for idx, slide in enumerate(prs.slides, start=1):
        lines = []
        for shape in slide.shapes:
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    t = "".join(run.text for run in para.runs).strip()
                    if t:
                        lines.append(t)
        if lines:
            for part in split_text("\n".join(lines)):
                chunks.append(Chunk(text=part, kind="text", source=path.name, page=idx))
    return chunks


def parse_xlsx(path: Path, max_rows: int = 200) -> list[Chunk]:
    wb = load_workbook(str(path), read_only=True, data_only=True)
    chunks: list[Chunk] = []
    for ws in wb.worksheets:
        rows = []
        for i, row in enumerate(ws.iter_rows(values_only=True)):
            if i >= max_rows:
                rows.append(["（表格过长，已截断）"])
                break
            if any(v is not None for v in row):
                cells = ["" if v is None else str(v).replace("\n", " ") for v in row]
                rows.append("| " + " | ".join(cells) + " |")
        if rows:
            header_sep = "| " + " | ".join(["---"] * (rows[0].count("|") - 1)) + " |"
            chunks.append(Chunk(
                text="\n".join([rows[0], header_sep, *rows[1:]]),
                kind="table",
                source=path.name,
                meta={"sheet": ws.title},
            ))
    wb.close()
    return chunks
